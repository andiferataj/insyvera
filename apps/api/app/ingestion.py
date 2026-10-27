import csv
import io
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from sqlalchemy import delete, select, text
from .config import settings
from .db import Chunk, Document, Version, transaction
from .storage import storage

SUPPORTED = {"pdf", "docx", "xlsx", "csv", "md", "txt"}

def numeric(value):
    from .sandbox_runner import number
    parsed = number(value)
    return float(parsed) if parsed is not None else None

def clean(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, bool, int, float)):
        return value
    return str(value)

def validate(content, filename):
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in SUPPORTED:
        raise ValueError("Supported files: PDF, DOCX, XLSX, CSV, Markdown and TXT.")
    if not content or len(content) > settings.max_upload:
        raise ValueError("Files must contain data and be at most 20 MB.")
    if ext == "pdf" and not content.startswith(b"%PDF-"):
        raise ValueError("File is not a valid PDF.")
    if ext in {"xlsx", "docx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 5000 or sum(v.file_size for v in infos) > 80 * 1024 * 1024:
                    raise ValueError("Expanded document exceeds the safety limit.")
                if any(v.flag_bits & 1 for v in infos):
                    raise ValueError("Encrypted files are not supported.")
                names = archive.namelist()
                marker = "xl/workbook.xml" if ext == "xlsx" else "word/document.xml"
                if marker not in names or any("vbaProject" in n for n in names):
                    raise ValueError("Invalid or macro-enabled Office file.")
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid Office archive.") from exc
    if ext in {"csv", "md", "txt"}:
        decoded = content.decode("utf-8-sig")
        if "\x00" in decoded:
            raise ValueError("Binary content is not valid text.")
    return ext

def profile(columns, rows):
    output = []
    for column in columns:
        values = [r.get(column) for r in rows]
        present = [v for v in values if v is not None and str(v).strip()]
        numbers = [numeric(v) for v in present]
        valid = [v for v in numbers if v is not None]
        kind = "number" if present and len(valid) == len(present) else "mixed" if valid else "text"
        stats = {}
        if valid:
            import numpy as np
            a = np.array(valid)
            q1, q3 = np.quantile(a, [.25, .75])
            iqr = q3 - q1
            stats = {"min": float(a.min()), "max": float(a.max()), "mean": float(a.mean()),
                "median": float(np.median(a)), "std": float(a.std()),
                "sum": str(sum((Decimal(str(v)) for v in valid), Decimal(0))),
                "anomalies": int(sum(v < q1 - 1.5 * iqr or v > q3 + 1.5 * iqr for v in valid))}
        output.append({"name": column, "type": kind, "missing": len(values) - len(present),
            "unique": len({str(v) for v in present}), "numeric_count": len(valid), **stats})
    duplicates = len(rows) - len({tuple(str(r.get(c)) for c in columns) for r in rows})
    return {"row_count": len(rows), "column_count": len(columns),
        "missing_cells": sum(c["missing"] for c in output), "duplicate_rows": duplicates, "columns": output}

def dataset(name, matrix, start_row=1, formulas=None, merged=None):
    if not matrix:
        return {"name": name, "columns": [], "rows": [], "row_numbers": [], "profile": profile([], []), "warnings": []}
    while matrix and not any(v is not None and str(v).strip() for v in matrix[0]):
        matrix = matrix[1:]
        start_row += 1
    if not matrix:
        return dataset(name, [])
    seen = set()
    columns = []
    for i, value in enumerate(matrix[0]):
        label = str(value).strip() if value is not None else f"Column {i + 1}"
        original = label
        suffix = 2
        while label in seen:
            label = f"{original} ({suffix})"
            suffix += 1
        seen.add(label)
        columns.append(label)
    rows, row_numbers = [], []
    for n, values in enumerate(matrix[1:], start_row + 1):
        if any(v is not None and str(v).strip() for v in values):
            rows.append({c: clean(values[i]) if i < len(values) else None for i, c in enumerate(columns)})
            row_numbers.append(n)
    warnings = []
    if formulas and any(f["value"] is None for f in formulas):
        warnings.append("Some formulas have no cached values. Recalculate in Excel or LibreOffice and upload a new version.")
    if merged:
        warnings.append("Merged cells are retained as source metadata; blank cells are not filled implicitly.")
    return {"name": name, "columns": columns, "rows": rows, "row_numbers": row_numbers,
        "range": f"A{start_row}:{excel_col(len(columns))}{max(row_numbers, default=start_row)}",
        "profile": profile(columns, rows), "formulas": formulas or [], "merged": merged or [], "warnings": warnings}

def excel_col(n):
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out or "A"

def parse(content, filename):
    ext = validate(content, filename)
    units, tables, warnings = [], [], []
    if ext == "pdf":
        import pymupdf
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            if pdf.is_encrypted:
                raise ValueError("Encrypted PDF files are not supported.")
            if len(pdf) > 1000:
                raise ValueError("PDF exceeds 1000 pages.")
            for page in pdf:
                value = page.get_text(sort=True).strip()
                units.append({"text": value, "location": {"page": page.number + 1}})
                if not value:
                    warnings.append(f"Page {page.number + 1} has no extractable text. OCR is required.")
                for i, table in enumerate(page.find_tables().tables):
                    tables.append(dataset(f"Page {page.number + 1} table {i + 1}", table.extract()))
    elif ext == "docx":
        from docx import Document as WordDocument
        word = WordDocument(io.BytesIO(content))
        for i, paragraph in enumerate(word.paragraphs):
            if paragraph.text.strip():
                units.append({"text": paragraph.text, "location": {"paragraph": i + 1}})
        for i, table in enumerate(word.tables):
            tables.append(dataset(f"Table {i + 1}", [[c.text for c in row.cells] for row in table.rows]))
    elif ext == "xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True, keep_links=False)
        if sum(s.max_row * s.max_column for s in workbook) > settings.max_cells:
            raise ValueError("Spreadsheet exceeds 200,000 cells.")
        for sheet in workbook:
            values, formulas = [], []
            for row in sheet.iter_rows():
                values.append([])
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        value = cached[sheet.title][cell.coordinate].value
                        formulas.append({"cell": cell.coordinate, "formula": cell.value, "value": clean(value)})
                    values[-1].append(clean(value))
            tables.append(dataset(sheet.title, values, formulas=formulas, merged=[str(r) for r in sheet.merged_cells.ranges]))
        workbook.close()
        cached.close()
    elif ext == "csv":
        value = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(value[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        matrix = list(csv.reader(io.StringIO(value), dialect))
        if sum(len(r) for r in matrix) > settings.max_cells:
            raise ValueError("Dataset exceeds 200,000 cells.")
        tables.append(dataset("Data", matrix))
    else:
        value = content.decode("utf-8-sig")
        for i, paragraph in enumerate(re.split(r"\n\s*\n", value)):
            if paragraph.strip():
                units.append({"text": paragraph.strip(), "location": {"section": i + 1}})
    for table in tables:
        warnings.extend(table["warnings"])
        for row, n in zip(table["rows"], table["row_numbers"]):
            units.append({"text": f'{table["name"]} | ' + " | ".join(f"{k}: {v}" for k, v in row.items()),
                "location": {"sheet": table["name"], "range": f'A{n}:{excel_col(len(table["columns"]))}{n}', "row": n}})
    return {"units": units, "datasets": tables, "warnings": sorted(set(warnings)), "parser": "insyvera-1"}

def ingest(version_id):
    # One version is serialized by a row lock in PostgreSQL. Replacement deletes old chunks atomically.
    with transaction() as db:
        version = db.scalar(select(Version).where(Version.id == version_id).with_for_update())
        if not version:
            return
        if version.status == "ready":
            return
        document = db.get(Document, version.document_id)
        try:
            parsed = parse(storage.get(version.key), document.name)
            db.execute(delete(Chunk).where(Chunk.version_id == version.id))
            count = 0
            for unit in parsed["units"]:
                for offset in range(0, max(1, len(unit["text"])), 1100):
                    value = unit["text"][offset:offset + 1250]
                    if not value.strip():
                        continue
                    count += 1
                    if count > settings.max_chunks:
                        raise ValueError("Document exceeds 10,000 searchable chunks.")
                    chunk = Chunk(version_id=version.id, workspace_id=document.workspace_id,
                        text=value, location={**unit["location"], "offset": offset})
                    if settings.embedding_provider != "lexical":
                        from .retrieval import embed
                        chunk.embedding = embed(value)
                        chunk.embedding_model = settings.embedding_model
                        chunk.dimensions = len(chunk.embedding)
                    db.add(chunk)
                    db.flush()
                    if chunk.embedding and db.bind.dialect.name == "postgresql":
                        db.execute(text("UPDATE document_chunks SET vector_embedding = CAST(:v AS vector) WHERE id=:id"),
                            {"v": str(chunk.embedding), "id": chunk.id})
            version.parsed = parsed
            version.status = "ready"
            version.error = None
        except Exception as exc:
            db.execute(delete(Chunk).where(Chunk.version_id == version.id))
            version.status = "failed"
            version.error = str(exc)[:500]
            version.parsed = {}
