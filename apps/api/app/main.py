import asyncio
import csv
import hashlib
import io
import json
import logging
import re
import time
from contextlib import asynccontextmanager
from typing import Literal
from fastapi import BackgroundTasks, Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import delete, func, select, text
from sqlalchemy.exc import IntegrityError
from .analytics import latest_sources, run_analysis
from .config import settings
from .db import (Analysis, Approval, Audit, Collection, Document, Evaluation, Member, Report,
    ReportVersion, Run, Session, Step, StorageDeletion, User, Version, Workspace, audit, initialize, transaction)
from .demo import fixtures
from .evaluation import evaluate
from .ingestion import validate
from .jobs import dispatch
from .reports import draft, pdf
from .retrieval import resolve, search
from .security import (authorize, create_session, current_user, get_db, limiter,
    password_hash, token_hash, verify_password)
from .storage import storage
from .tools import calculator, public_url, retrieve_web

logger = logging.getLogger("insyvera")
logging.basicConfig(level=logging.INFO, format="%(message)s")

@asynccontextmanager
async def lifespan(app):
    initialize()
    storage.ready()
    if __import__("os").getenv("OTEL_EXPORTER_OTLP_ENDPOINT"):
        from opentelemetry import trace
        from opentelemetry.sdk.trace import TracerProvider
        from opentelemetry.sdk.trace.export import BatchSpanProcessor
        from opentelemetry.exporter.otlp.proto.http.trace_exporter import OTLPSpanExporter
        provider = TracerProvider()
        provider.add_span_processor(BatchSpanProcessor(OTLPSpanExporter()))
        trace.set_tracer_provider(provider)
    yield

app = FastAPI(title="Insyvera", version="0.1.0", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=settings.origins.split(","),
    allow_credentials=True, allow_methods=["GET","POST","PATCH","DELETE"], allow_headers=["Content-Type"])

@app.middleware("http")
async def boundary(request: Request, call_next):
    start = time.monotonic()
    if int(request.headers.get("content-length", "0")) > 21 * 1024 * 1024:
        return JSONResponse({"detail":"Request exceeds 21 MB."}, status_code=413)
    if request.headers.get("transfer-encoding", "").lower() == "chunked" and request.method in {"POST","PUT","PATCH"}:
        return JSONResponse({"detail":"Supply a Content-Length for uploads and JSON requests."}, status_code=411)
    origin = request.headers.get("origin")
    if request.method in {"POST","PATCH","DELETE","PUT"} and origin and origin not in settings.origins.split(","):
        return JSONResponse({"detail": "Origin is not allowed."}, status_code=403)
    from opentelemetry import trace
    with trace.get_tracer("insyvera.api").start_as_current_span(request.method + " " + request.url.path) as span:
        response = await call_next(request)
        span.set_attribute("http.response.status_code", response.status_code)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Cache-Control"] = "no-store"
    logger.info(json.dumps({"method": request.method, "path": request.url.path,
        "status": response.status_code, "duration_ms": round((time.monotonic()-start)*1000)}))
    return response

@app.exception_handler(ValueError)
async def value_error(request, exc):
    return JSONResponse({"detail": str(exc)[:500]}, status_code=422)

@app.exception_handler(IntegrityError)
async def conflict(request, exc):
    return JSONResponse({"detail": "This record already exists or changed. Refresh and retry."}, status_code=409)

@app.get("/api/health")
def health():
    return {"status": "ok", "version": "0.1.0"}

@app.get("/api/ready")
def ready():
    checks = {}
    try:
        with transaction() as db:
            db.execute(text("SELECT 1"))
            if db.bind.dialect.name == "postgresql":
                db.execute(text("SELECT '[1,2,3]'::vector"))
        checks["database"] = "ok"
        storage.ready()
        checks["storage"] = "ok"
        if settings.redis_url:
            import redis
            redis.Redis.from_url(settings.redis_url, socket_connect_timeout=3).ping()
            checks["queue"] = "ok"
        else:
            checks["queue"] = "local-background"
    except Exception:
        return JSONResponse({"status": "unavailable", "checks": checks}, status_code=503)
    return {"status": "ready", "checks": checks, "model_provider": settings.provider,
        "retrieval": "lexical" if settings.embedding_provider == "lexical" else "hybrid"}

class Credentials(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=10, max_length=128)
    name: str = Field(default="Researcher", min_length=1, max_length=100)
    @field_validator("email")
    @classmethod
    def email_valid(cls, value):
        value = value.strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
            raise ValueError("Enter a valid email address.")
        return value

class NameInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    @field_validator("name")
    @classmethod
    def trim(cls, value):
        if not value.strip():
            raise ValueError("Name cannot be blank.")
        return value.strip()

@app.post("/api/auth/register", status_code=201)
def register(body: Credentials, request: Request, response: Response, db=Depends(get_db)):
    limiter.check("register:" + (request.client.host if request.client else "unknown"), 10)
    if db.scalar(select(User).where(User.email == body.email)):
        raise HTTPException(409, "An account with this email already exists.")
    user = User(email=body.email, name=body.name.strip(), password=password_hash(body.password))
    db.add(user)
    db.flush()
    workspace = Workspace(name="My workspace")
    db.add(workspace)
    db.flush()
    db.add(Member(workspace_id=workspace.id, user_id=user.id, role="owner"))
    create_session(db, user, response)
    audit(db, "account.created", user.id, workspace.id)
    db.commit()
    return {"id": user.id, "name": user.name, "email": user.email}

@app.post("/api/auth/login")
def login(body: Credentials, request: Request, response: Response, db=Depends(get_db)):
    limiter.check("login:" + (request.client.host if request.client else "unknown"))
    user = db.scalar(select(User).where(User.email == body.email))
    if not user or not verify_password(body.password, user.password):
        raise HTTPException(401, "Email or password is incorrect.")
    create_session(db, user, response)
    audit(db, "account.signed_in", user.id)
    db.commit()
    return {"id": user.id, "name": user.name, "email": user.email}

@app.post("/api/auth/logout")
def logout(request: Request, response: Response, user=Depends(current_user), db=Depends(get_db)):
    db.execute(delete(Session).where(Session.token == token_hash(request.cookies.get("insyvera_session", ""))))
    db.commit()
    response.delete_cookie("insyvera_session", path="/")
    return {"ok": True}

@app.get("/api/me")
def me(user=Depends(current_user)):
    return {"id": user.id, "name": user.name, "email": user.email}

@app.get("/api/workspaces")
def workspaces(user=Depends(current_user), db=Depends(get_db)):
    rows = db.execute(select(Workspace, Member.role).join(Member).where(Member.user_id == user.id).order_by(Workspace.created_at)).all()
    return [{"id": w.id, "name": w.name, "role": role, "settings": w.settings} for w,role in rows]

@app.post("/api/workspaces", status_code=201)
def create_workspace(body: NameInput, user=Depends(current_user), db=Depends(get_db)):
    w = Workspace(name=body.name)
    db.add(w)
    db.flush()
    db.add(Member(workspace_id=w.id, user_id=user.id))
    audit(db, "workspace.created", user.id, w.id)
    db.commit()
    return {"id": w.id, "name": w.name, "role": "owner", "settings": {}}

@app.patch("/api/workspaces/{workspace}")
def rename_workspace(workspace: str, body: NameInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.get(Workspace, workspace)
    item.name = body.name
    audit(db, "workspace.renamed", user.id, workspace)
    db.commit()
    return {"id": item.id, "name": item.name}

@app.get("/api/workspaces/{workspace}/collections")
def collections(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"id": c.id, "name": c.name} for c in db.scalars(select(Collection).where(Collection.workspace_id == workspace))]

@app.post("/api/workspaces/{workspace}/collections", status_code=201)
def create_collection(workspace: str, body: NameInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = Collection(workspace_id=workspace, name=body.name)
    db.add(item)
    db.commit()
    return {"id": item.id, "name": item.name}

def document_item(doc, versions):
    return {"id": doc.id, "name": doc.name, "type": doc.type, "collection_id": doc.collection_id,
        "created_at": doc.created_at.isoformat(), "versions": [{"id": v.id, "number": v.number,
        "size": v.size, "hash": v.hash, "status": v.status, "error": v.error,
        "warnings": v.parsed.get("warnings", []), "datasets": len(v.parsed.get("datasets", [])),
        "created_at": v.created_at.isoformat()} for v in versions]}

@app.get("/api/workspaces/{workspace}/documents")
def documents(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    docs = db.scalars(select(Document).where(Document.workspace_id == workspace).order_by(Document.created_at.desc())).all()
    return [document_item(d, db.scalars(select(Version).where(Version.document_id == d.id).order_by(Version.number.desc())).all()) for d in docs]

def store_upload(db, workspace, user, content, name, document_id=None, collection_id=None):
    name = name.replace("\\", "/").split("/")[-1][:240]
    if not name:
        raise ValueError("A filename is required.")
    ext = validate(content, name)
    digest = hashlib.sha256(content).hexdigest()
    if document_id:
        doc = db.scalar(select(Document).where(Document.id == document_id, Document.workspace_id == workspace).with_for_update())
        if not doc:
            raise HTTPException(404, "File not found.")
        if ext != doc.type:
            raise ValueError("A new version must use the same file type.")
    else:
        existing = db.execute(select(Document, Version).join(Version).where(
            Document.workspace_id == workspace, Document.name == name, Version.hash == digest)).first()
        if existing:
            return existing[0], existing[1], False
        doc = Document(workspace_id=workspace, name=name, type=ext, collection_id=collection_id)
        db.add(doc)
        db.flush()
    if collection_id and not db.scalar(select(Collection).where(Collection.id == collection_id, Collection.workspace_id == workspace)):
        raise HTTPException(404, "Collection not found.")
    number = (db.scalar(select(func.max(Version.number)).where(Version.document_id == doc.id)) or 0) + 1
    from .db import uid
    version_id = uid()
    key = f"{workspace}/{doc.id}/{version_id}/original.{ext}"
    storage.put(key, content)
    version = Version(id=version_id, document_id=doc.id, number=number, hash=digest, size=len(content),
        key=key, status="queued")
    db.add(version)
    db.flush()
    audit(db, "document.uploaded", user.id, workspace, document_id=doc.id, version_id=version.id, sha256=digest)
    return doc, version, True

@app.post("/api/workspaces/{workspace}/documents", status_code=201)
async def upload(workspace: str, background: BackgroundTasks, file: UploadFile = File(...),
    collection_id: str | None = None, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    limiter.check("upload:" + user.id, 60)
    content = await file.read(settings.max_upload+1)
    doc, version, created = store_upload(db, workspace, user, content, file.filename or "", collection_id=collection_id)
    db.commit()
    if created:
        background.add_task(dispatch, "ingest", version.id)
    return document_item(doc, [version])

@app.get("/api/workspaces/{workspace}/documents/{document_id}")
def preview(workspace: str, document_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    doc = db.scalar(select(Document).where(Document.workspace_id == workspace, Document.id == document_id))
    if not doc:
        raise HTTPException(404, "File not found.")
    versions = db.scalars(select(Version).where(Version.document_id == doc.id).order_by(Version.number.desc())).all()
    return {**document_item(doc, versions), "parsed_versions": {v.id: v.parsed for v in versions}}

@app.post("/api/workspaces/{workspace}/documents/{document_id}/versions", status_code=201)
async def add_version(workspace: str, document_id: str, background: BackgroundTasks, file: UploadFile = File(...),
    user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    doc, version, _ = store_upload(db, workspace, user, await file.read(settings.max_upload+1), file.filename or "", document_id)
    db.commit()
    background.add_task(dispatch, "ingest", version.id)
    return document_item(doc, [version])

def scoped_version(db, workspace, version_id):
    row = db.execute(select(Document, Version).join(Version).where(Document.workspace_id == workspace, Version.id == version_id)).first()
    if not row:
        raise HTTPException(404, "Source version not found.")
    return row

@app.get("/api/workspaces/{workspace}/versions/{version_id}/content")
def original(workspace: str, version_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    doc, version = scoped_version(db, workspace, version_id)
    mime = {"pdf": "application/pdf", "txt": "text/plain", "md": "text/plain",
        "csv": "text/csv", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}[doc.type]
    return Response(storage.get(version.key), media_type=mime,
        headers={"Content-Disposition": "inline" if doc.type == "pdf" else "attachment"})

@app.post("/api/workspaces/{workspace}/versions/{version_id}/reprocess")
def reprocess(workspace: str, version_id: str, background: BackgroundTasks, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    _, version = scoped_version(db, workspace, version_id)
    if version.status in {"processing", "queued"}:
        raise HTTPException(409, "This source is already queued.")
    version.status = "queued"
    version.error = None
    audit(db, "document.reprocessed", user.id, workspace, version_id=version_id)
    db.commit()
    background.add_task(dispatch, "ingest", version_id)
    return {"status": "queued"}

@app.delete("/api/workspaces/{workspace}/documents/{document_id}")
def delete_document(workspace: str, document_id: str, background: BackgroundTasks, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    doc = db.scalar(select(Document).where(Document.id == document_id, Document.workspace_id == workspace).with_for_update())
    if not doc:
        raise HTTPException(404, "File not found.")
    versions = db.scalars(select(Version).where(Version.document_id == document_id)).all()
    cleanup = []
    for v in versions:
        task = StorageDeletion(key=v.key)
        db.add(task)
        db.flush()
        cleanup.append(task.id)
    db.delete(doc)
    audit(db, "document.deleted", user.id, workspace, document_id=document_id)
    db.commit()
    from .gc import purge
    for identifier in cleanup:
        background.add_task(purge, identifier)
    return {"ok": True, "purge_status": "scheduled"}

@app.post("/api/workspaces/{workspace}/demo")
def seed_demo(workspace: str, background: BackgroundTasks, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    queued = []
    for name, content in fixtures().items():
        if db.scalar(select(Document.id).where(Document.workspace_id == workspace, Document.name == name)):
            continue
        _, version, created = store_upload(db, workspace, user, content, name)
        if created:
            queued.append(version.id)
    db.commit()
    for version_id in queued:
        background.add_task(dispatch, "ingest", version_id)
    return {"created": len(queued), "synthetic": True}

@app.get("/api/workspaces/{workspace}/search")
def find(workspace: str, q: str = Query(min_length=1, max_length=1000), limit: int = Query(default=8, ge=1, le=30),
    collection_id: str | None = None, document_id: str | None = None, version_id: str | None = None,
    user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return search(db, workspace, q, limit, [document_id] if document_id else None, collection_id, version_id)

@app.get("/api/workspaces/{workspace}/sources/{chunk_id}")
def source(workspace: str, chunk_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    result = resolve(db, workspace, chunk_id)
    if not result:
        raise HTTPException(404, "Evidence source no longer exists.")
    return result

class ResearchInput(BaseModel):
    question: str = Field(min_length=3, max_length=4000)
    mode: Literal["research", "analysis", "investigation"] = "research"
    document_ids: list[str] | None = Field(default=None, max_length=50)
    provider: Literal["extractive", "ollama", "openai"] | None = None
    model: str | None = Field(default=None, max_length=120)

@app.post("/api/workspaces/{workspace}/runs", status_code=201)
def start_run(workspace: str, body: ResearchInput, background: BackgroundTasks, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    limiter.check("research:" + user.id, 30)
    if body.document_ids:
        for identifier in body.document_ids:
            if not db.scalar(select(Document.id).where(Document.id == identifier, Document.workspace_id == workspace)):
                raise HTTPException(404, "Selected source not found.")
    provider = body.provider or settings.provider
    if provider == "openai" and not settings.openai_key:
        raise HTTPException(422, "Configure an API key before selecting the hosted provider.")
    run = Run(workspace_id=workspace, question=body.question, mode=body.mode,
        config={"document_ids": body.document_ids, "provider": provider, "model": "extractive-v1" if provider == "extractive" else body.model or settings.model})
    db.add(run)
    db.flush()
    audit(db, "research.started", user.id, workspace, run_id=run.id)
    db.commit()
    background.add_task(dispatch, "research", run.id)
    return {"id": run.id, "status": "queued"}

def scoped_run(db, workspace, run_id):
    item = db.scalar(select(Run).where(Run.workspace_id == workspace, Run.id == run_id))
    if not item:
        raise HTTPException(404, "Research run not found.")
    return item

def run_json(run, db=None):
    result = dict(run.result or {})
    if db and result:
        deleted = []
        for s in result.get("sources", []):
            if not resolve(db, run.workspace_id, s["id"]):
                deleted.append(s["id"])
        result["unavailable_source_ids"] = deleted
        if deleted:
            result["warnings"] = result.get("warnings", []) + ["Some cited sources were deleted or reprocessed. Their archived quotations cannot currently be resolved."]
    return {"id": run.id, "question": run.question, "mode": run.mode, "status": run.status,
        "config": run.config, "result": result, "error": run.error, "created_at": run.created_at.isoformat()}

@app.get("/api/workspaces/{workspace}/runs")
def runs(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [run_json(r) for r in db.scalars(select(Run).where(Run.workspace_id == workspace).order_by(Run.created_at.desc()).limit(100))]

@app.get("/api/workspaces/{workspace}/runs/{run_id}")
def run_detail(workspace: str, run_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    run = scoped_run(db, workspace, run_id)
    steps = db.scalars(select(Step).where(Step.run_id == run_id).order_by(Step.created_at)).all()
    return {**run_json(run, db), "steps": [{"id": s.id, "node": s.node, "status": s.status, "detail": s.detail,
        "duration_ms": s.duration_ms, "created_at": s.created_at.isoformat()} for s in steps]}

@app.get("/api/workspaces/{workspace}/runs/{run_id}/events")
def run_events(workspace: str, run_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    scoped_run(db, workspace, run_id)
    async def stream():
        for _ in range(210):
            with transaction() as current:
                run = scoped_run(current, workspace, run_id)
                value = run_json(run)
                yield "data: " + json.dumps(value) + "\n\n"
                if run.status in {"completed", "failed", "cancelled"}:
                    return
            await asyncio.sleep(1)
    return StreamingResponse(stream(), media_type="text/event-stream")

@app.post("/api/workspaces/{workspace}/runs/{run_id}/cancel")
def cancel_run(workspace: str, run_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    run = scoped_run(db, workspace, run_id)
    if run.status not in {"queued", "running"}:
        raise HTTPException(409, "Only queued or running research can be cancelled.")
    run.status = "cancelled"
    audit(db, "research.cancelled", user.id, workspace, run_id=run_id)
    db.commit()
    return {"status": "cancelled"}

@app.post("/api/workspaces/{workspace}/runs/{run_id}/resume")
def resume_run(workspace: str, run_id: str, background: BackgroundTasks, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    run = scoped_run(db, workspace, run_id)
    if run.status not in {"failed", "cancelled"}:
        from .db import now
        if run.status != "running" or (now() - run.updated_at.replace(tzinfo=now().tzinfo)).total_seconds() < 240:
            raise HTTPException(409, "This run cannot be resumed.")
    from .locks import BusyRun, run_lock
    try:
        with run_lock(run_id):
            pass
    except BusyRun as exc:
        raise HTTPException(409, str(exc)) from exc
    run.status = "queued"
    db.commit()
    background.add_task(dispatch, "research", run_id, True)
    return {"status": "queued"}

class AnalysisInput(BaseModel):
    version_id: str
    table_index: int = Field(default=0, ge=0)
    operation: Literal["aggregate", "statistics", "filter", "join", "correlation", "clean", "trend"] = "aggregate"
    column: str | None = None
    group_by: str | None = None
    second_column: str | None = None
    rule: dict | None = None
    options: dict | None = None
    right_version_id: str | None = None
    right_table_index: int = Field(default=0, ge=0)
    key: str | None = None

@app.post("/api/workspaces/{workspace}/analyses", status_code=201)
def analyze(workspace: str, body: AnalysisInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    parameters = body.model_dump(exclude={"version_id", "table_index", "operation", "right_version_id", "right_table_index"}, exclude_none=True)
    if body.operation == "join":
        _, right = scoped_version(db, workspace, body.right_version_id or "")
        tables = right.parsed.get("datasets", [])
        if body.right_table_index >= len(tables):
            raise ValueError("Right dataset index not found.")
        parameters["right_rows"] = tables[body.right_table_index]["rows"]
        parameters["right_source"] = {"version_id": right.id, "table_index": body.right_table_index, "hash": right.hash}
    result = run_analysis(db, workspace, body.version_id, body.table_index, body.operation, **parameters)
    audit(db, "analysis.executed", user.id, workspace, analysis_id=result["id"])
    db.commit()
    return result

@app.get("/api/workspaces/{workspace}/datasets")
def datasets(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"document_id": d.id, "version_id": v.id, "document_name": d.name, "table_index": i, **t}
        for d,v in latest_sources(db, workspace) for i,t in enumerate(v.parsed.get("datasets", []))]

@app.get("/api/workspaces/{workspace}/analyses")
def analyses(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"id": a.id, "definition": a.definition, "result": a.result, "created_at": a.created_at.isoformat()}
        for a in db.scalars(select(Analysis).where(Analysis.workspace_id == workspace).order_by(Analysis.created_at.desc()).limit(100))]

@app.get("/api/workspaces/{workspace}/analyses/{analysis_id}/export")
def export_analysis(workspace: str, analysis_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.scalar(select(Analysis).where(Analysis.id == analysis_id, Analysis.workspace_id == workspace))
    if not item:
        raise HTTPException(404, "Analysis not found.")
    result = item.result["result"]
    rows = result.get("rows", result.get("groups", [result]))
    output = io.StringIO()
    fields = list(dict.fromkeys(k for row in rows for k in row))
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for row in rows:
        writer.writerow({k: "'" + str(v) if str(v).startswith(("=", "+", "-", "@", "\t", "\r")) else v for k,v in row.items()})
    return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="analysis.csv"'})

class ReportInput(BaseModel):
    title: str = Field(default="Investigation report", min_length=1, max_length=200)
    run_id: str | None = None
    content: str | None = Field(default=None, max_length=100000)
    expected_version: int | None = None

def report_json(db, report):
    versions = db.scalars(select(ReportVersion).where(ReportVersion.report_id == report.id).order_by(ReportVersion.number.desc())).all()
    return {"id": report.id, "title": report.title, "run_id": report.run_id,
        "created_at": report.created_at.isoformat(), "versions": [{"id": v.id, "number": v.number,
            "content": v.content, "created_at": v.created_at.isoformat()} for v in versions]}

@app.get("/api/workspaces/{workspace}/reports")
def reports(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [report_json(db,r) for r in db.scalars(select(Report).where(Report.workspace_id == workspace).order_by(Report.created_at.desc()))]

@app.post("/api/workspaces/{workspace}/reports", status_code=201)
def create_report(workspace: str, body: ReportInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    run = scoped_run(db, workspace, body.run_id) if body.run_id else None
    if run and run.status != "completed":
        raise HTTPException(409, "Complete the investigation before creating a report.")
    item = Report(workspace_id=workspace, title=body.title, run_id=body.run_id)
    db.add(item)
    db.flush()
    db.add(ReportVersion(report_id=item.id, number=1, content=body.content or (draft(run) if run else "# " + body.title)))
    audit(db, "report.created", user.id, workspace, report_id=item.id)
    db.commit()
    return report_json(db, item)

@app.patch("/api/workspaces/{workspace}/reports/{report_id}")
def edit_report(workspace: str, report_id: str, body: ReportInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.scalar(select(Report).where(Report.id == report_id, Report.workspace_id == workspace).with_for_update())
    if not item:
        raise HTTPException(404, "Report not found.")
    latest = db.scalar(select(func.max(ReportVersion.number)).where(ReportVersion.report_id == report_id))
    if body.expected_version != latest:
        raise HTTPException(409, "Report has changed. Refresh before saving.")
    item.title = body.title
    db.add(ReportVersion(report_id=report_id, number=latest+1, content=body.content or ""))
    audit(db, "report.version_saved", user.id, workspace, report_id=report_id, version=latest+1)
    db.commit()
    return report_json(db, item)

@app.get("/api/workspaces/{workspace}/reports/{report_id}/export")
def export_report(workspace: str, report_id: str, format: Literal["md", "pdf", "docx"] = "md",
    version: int | None = None, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.scalar(select(Report).where(Report.id == report_id, Report.workspace_id == workspace))
    if not item:
        raise HTTPException(404, "Report not found.")
    query = select(ReportVersion).where(ReportVersion.report_id == item.id)
    if version:
        query = query.where(ReportVersion.number == version)
    result = db.scalar(query.order_by(ReportVersion.number.desc()))
    if not result:
        raise HTTPException(404, "Report version not found.")
    if format == "pdf":
        content, mime = pdf(item.title, result.content), "application/pdf"
    elif format == "docx":
        from docx import Document as WordDocument
        document = WordDocument()
        for line in result.content.splitlines():
            if line.startswith("#"):
                document.add_heading(line.lstrip("# "), min(len(line)-len(line.lstrip("#")), 3))
            else:
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        content, mime = output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        content, mime = result.content.encode(), "text/markdown"
    return Response(content, media_type=mime, headers={"Content-Disposition": f'attachment; filename="insyvera-report.{format}"'})

class EvaluationInput(BaseModel):
    cases: list[dict] | None = Field(default=None, max_length=30)
    provider: Literal["extractive", "ollama", "openai"] = "extractive"
    model: str = Field(default="extractive", max_length=120)

@app.post("/api/workspaces/{workspace}/evaluations", status_code=201)
def start_evaluation(workspace: str, body: EvaluationInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    limiter.check("eval:" + user.id, 5)
    return evaluate(workspace, body.cases, body.provider, body.model)

@app.get("/api/workspaces/{workspace}/evaluations")
def evaluations(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"id": e.id, "config": e.config, "scores": e.scores, "created_at": e.created_at.isoformat()}
        for e in db.scalars(select(Evaluation).where(Evaluation.workspace_id == workspace).order_by(Evaluation.created_at.desc()).limit(50))]

class ApprovalInput(BaseModel):
    url: str = Field(max_length=2000)

@app.post("/api/workspaces/{workspace}/approvals", status_code=201)
def request_approval(workspace: str, body: ApprovalInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    if not settings.external_web:
        raise HTTPException(403, "External web retrieval is disabled in server settings.")
    public_url(body.url)
    item = Approval(workspace_id=workspace, user_id=user.id, tool="web.retrieve", target=body.url)
    db.add(item)
    audit(db, "tool.approval_requested", user.id, workspace, target=body.url)
    db.commit()
    return {"id": item.id, "tool": item.tool, "target": item.target, "status": item.status}

@app.get("/api/workspaces/{workspace}/approvals")
def approvals(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"id": a.id, "tool": a.tool, "target": a.target, "status": a.status, "result": a.result}
        for a in db.scalars(select(Approval).where(Approval.workspace_id == workspace).order_by(Approval.created_at.desc()))]

@app.post("/api/workspaces/{workspace}/approvals/{approval_id}/{decision}")
def approve_tool(workspace: str, approval_id: str, decision: Literal["approve", "deny"],
    user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.scalar(select(Approval).where(Approval.id == approval_id, Approval.workspace_id == workspace).with_for_update())
    if not item or item.status != "pending" or item.user_id != user.id:
        raise HTTPException(404, "Pending approval not found.")
    if decision == "approve":
        if not settings.external_web:
            raise HTTPException(403, "External retrieval is disabled.")
        item.result = retrieve_web(item.target)
        item.status = "completed"
    else:
        item.status = "denied"
    audit(db, "tool." + item.status, user.id, workspace, approval_id=item.id)
    db.commit()
    return {"status": item.status, "result": item.result}

class CalculatorInput(BaseModel):
    expression: str = Field(max_length=200)

@app.post("/api/workspaces/{workspace}/calculator")
def calculate(workspace: str, body: CalculatorInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    value = calculator(body.expression)
    audit(db, "tool.calculator", user.id, workspace, expression=body.expression, result=value)
    db.commit()
    return {"result": value}

@app.get("/api/workspaces/{workspace}/audit")
def audit_history(workspace: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    return [{"id": a.id, "action": a.action, "detail": a.detail, "created_at": a.created_at.isoformat()}
        for a in db.scalars(select(Audit).where(Audit.workspace_id == workspace).order_by(Audit.created_at.desc()).limit(100))]

@app.get("/api/config")
def configuration(user=Depends(current_user)):
    return {"provider": settings.provider, "model": settings.model, "embedding_provider": settings.embedding_provider,
        "embedding_model": settings.embedding_model, "dimensions": settings.dimensions,
        "external_web": settings.external_web, "max_upload_mb": settings.max_upload//1024//1024,
        "execution": "Allowlisted operations only. Arbitrary generated Python is not accepted."}

class ChartInput(BaseModel):
    version_id: str
    table_index: int = Field(default=0, ge=0)
    kind: Literal["bar", "line", "histogram", "scatter"] = "bar"
    x: str
    y: str | None = None
    bins: int = Field(default=12, ge=2, le=60)

@app.post("/api/workspaces/{workspace}/charts", status_code=201)
def create_chart(workspace: str, body: ChartInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    from .visualization import chart_analysis
    result = chart_analysis(db, workspace, body.model_dump())
    db.commit()
    return result

@app.post("/api/workspaces/{workspace}/analyses/{analysis_id}/replay", status_code=201)
def replay_analysis(workspace: str, analysis_id: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    item = db.scalar(select(Analysis).where(Analysis.id == analysis_id, Analysis.workspace_id == workspace))
    if not item:
        raise HTTPException(404, "Analysis not found.")
    definition = dict(item.definition)
    if definition.get("operation") == "chart":
        from .visualization import chart_analysis
        result = chart_analysis(db, workspace, {k:v for k,v in definition.items() if k != "operation"})
    else:
        source_version = definition.pop("version_id")
        table_index = definition.pop("table_index")
        operation = definition.pop("operation")
        if "right_source" in definition:
            source = definition["right_source"]
            _, version = scoped_version(db, workspace, source["version_id"])
            definition["right_rows"] = version.parsed["datasets"][source["table_index"]]["rows"]
        result = run_analysis(db, workspace, source_version, table_index, operation, **definition)
    audit(db, "analysis.replayed", user.id, workspace, original_analysis_id=analysis_id, analysis_id=result["id"])
    db.commit()
    return result

@app.get("/api/workspaces/{workspace}/documents/{document_id}/compare")
def compare_versions(workspace: str, document_id: str, before: str, after: str, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    import difflib
    first_doc, first = scoped_version(db, workspace, before)
    second_doc, second = scoped_version(db, workspace, after)
    if first_doc.id != document_id or second_doc.id != document_id:
        raise HTTPException(404, "Versions do not belong to this source.")
    a = "\n".join(u["text"] for u in first.parsed.get("units", [])).splitlines()
    b = "\n".join(u["text"] for u in second.parsed.get("units", [])).splitlines()
    differences = list(difflib.unified_diff(a, b, fromfile=f"version-{first.number}", tofile=f"version-{second.number}", lineterm=""))
    return {"document_id": document_id, "before": {"id":first.id,"hash":first.hash,"number":first.number},
        "after": {"id":second.id,"hash":second.hash,"number":second.number},
        "diff": differences[:2000], "truncated": len(differences)>2000,
        "changed": first.hash != second.hash, "comparison": "extracted text and row representation"}

class DocumentCollectionInput(BaseModel):
    collection_id: str | None = None

@app.patch("/api/workspaces/{workspace}/documents/{document_id}")
def organize_document(workspace: str, document_id: str, body: DocumentCollectionInput, user=Depends(current_user), db=Depends(get_db)):
    authorize(db, workspace, user)
    document = db.scalar(select(Document).where(Document.workspace_id == workspace, Document.id == document_id))
    if not document:
        raise HTTPException(404, "Document not found.")
    if body.collection_id and not db.scalar(select(Collection).where(Collection.id == body.collection_id, Collection.workspace_id == workspace)):
        raise HTTPException(404, "Collection not found.")
    document.collection_id = body.collection_id
    audit(db, "document.organized", user.id, workspace, document_id=document_id, collection_id=body.collection_id)
    db.commit()
    return {"id": document.id, "collection_id": document.collection_id}
