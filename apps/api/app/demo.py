"""Synthetic, versioned acceptance fixtures. No real personal or financial data."""
import io
import json
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from docx import Document
from reportlab.pdfgen.canvas import Canvas

def workbook(sheet, rows):
    book = openpyxl.Workbook()
    book.properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    book.properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    ws = book.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="35664C")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 20
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()

def fixtures():
    taxes = workbook("Transactions", [
        ["Category", "Type", "Amount", "Period", "Currency"],
        ["Salary", "Income", 72000, "2026", "EUR"], ["Freelance", "Income", 12000, "2026", "EUR"],
        ["Equipment", "Expense", 9000, "2026", "EUR"], ["Travel", "Expense", 6000, "2026", "EUR"],
        ["Software", "Expense", 3000, "2026", "EUR"]])
    gifts = workbook("Gifts", [
        ["Recipient", "Category", "Amount", "Budget", "Date", "Currency"],
        ["Mira", "Electronics", 150, 120, "2026-12-01", "EUR"],
        ["Arben", "Books", 80, 80, "2026-12-02", "EUR"],
        ["Lina", "Home", 120, 100, "2026-12-04", "EUR"],
        ["Dion", "Games", 90, 80, "2026-12-05", "EUR"],
        ["Sara", None, 70, 60, "2026-12-05", "EUR"],
        ["Noel", "Books", 60, 60, "2026-12-07", "EUR"],
        ["Unassigned", "Pending", None, None, "2026-12-08", "EUR"]])
    pdf_bytes = io.BytesIO()
    canvas = Canvas(pdf_bytes, invariant=1)
    for y, line in enumerate(["Synthetic tax summary", "Period: 2026", "Currency: EUR", "Income: 84500",
        "Expenses: 18000", "Taxable income: 66500", "Filing deadline: April 15, 2027",
        "This fixture deliberately contains a discrepancy. It is not a tax document."]):
        canvas.drawString(60, 780-y*28, line)
    canvas.save()
    word = Document()
    word.add_heading("Research notes", 0)
    word.add_paragraph("Period: 2026")
    word.add_paragraph("Filing deadline: April 30, 2027")
    word.add_paragraph("The deadline conflicts with the synthetic PDF. Verify the authoritative source.")
    doc = io.BytesIO()
    word.save(doc)
    return {"taxes_2026.xlsx": taxes, "tax_report_2026.pdf": pdf_bytes.getvalue(),
        "christmas_gifts.xlsx": gifts, "research_notes.docx": doc.getvalue(),
        "quality_edges.csv": b"Item,Amount,Category\nAlpha,10,A\nBeta,20,B\nBeta,20,B\nGamma,,A\nOutlier,1000,C\n",
        "project_brief.md": b"# Insyvera demo\n\nAll fixtures are synthetic. Source files remain immutable.\n",
        "readme.txt": b"These records do not contain bank credentials or establish available cash."}

def write_fixtures(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    for name, content in fixtures().items():
        (directory / name).write_bytes(content)
    (directory / "expected.json").write_text(json.dumps({"income": 84000, "expenses": 18000,
        "taxable_income": 66000, "reported_income": 84500, "income_difference": 500,
        "gift_spending": 570, "gift_budget": 500, "over_budget": 70}, indent=2))
